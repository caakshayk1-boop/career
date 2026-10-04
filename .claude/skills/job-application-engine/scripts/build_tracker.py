#!/usr/bin/env python3
"""Build the job application tracker workbook.

Sheets: Summary, Applications, CV Details, Skipped Jobs.
Upload the result to Google Drive to get a live, editable Google Sheet.

Usage:
    python3 build_tracker.py --applications applications.json \
        --out outputs/Job_Application_Tracker.xlsx
"""
import argparse
import json
import os
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HEAD_FILL = PatternFill("solid", fgColor="1F2937")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
BAND = {
    "Strong": PatternFill("solid", fgColor="D1FAE5"),
    "Worth a shot": PatternFill("solid", fgColor="FEF3C7"),
    "Skip": PatternFill("solid", fgColor="FEE2E2"),
}
THIN = Border(*[Side(style="thin", color="E5E7EB")] * 4)


def _header(ws, headers, widths):
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for cell in ws[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", horizontal="left")
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22


def _summary(wb, data):
    ws = wb.create_sheet("Summary", 0)
    apps = data.get("applications", [])
    skipped = data.get("skipped", [])
    insights = data.get("insights", {})
    strong = sum(1 for a in apps if a.get("band") == "Strong")
    flags = sum(int(a.get("need_number_flags", 0)) for a in apps)

    ws["A1"] = "Job Application Report"
    ws["A1"].font = Font(bold=True, size=16)
    ws["A2"] = f"{data.get('candidate', '')}  ·  generated {data.get('generated', date.today().isoformat())}"
    ws["A2"].font = Font(size=10, color="6B7280")

    rows = [
        ("Jobs reviewed", len(apps) + len(skipped)),
        ("CVs generated", len(apps)),
        ("Strong matches (75+)", strong),
        ("Skipped as poor fit", len(skipped)),
        ("[NEED NUMBER] placeholders to fill before sending", flags),
        ("Most-requested skill missing from profile", insights.get("top_missing_skill", "—")),
        ("Strongest match", insights.get("best_match", "—")),
    ]
    for i, (label, value) in enumerate(rows, start=4):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value)
    ws.column_dimensions["A"].width = 48
    ws.column_dimensions["B"].width = 46

    ws.cell(row=len(rows) + 5, column=1,
            value="Fill every [NEED NUMBER] placeholder with a real figure before sending. "
                  "Never guess — an invented metric fails at interview.").font = Font(
        italic=True, color="B91C1C")


def _applications(wb, apps):
    ws = wb.create_sheet("Applications")
    _header(ws, ["#", "Company", "Role", "Location", "Salary", "Fit", "Band",
                 "Why", "CV File", "Job Link", "Status", "Date Applied",
                 "Follow Up", "Notes"],
            [5, 22, 28, 18, 16, 6, 13, 46, 34, 16, 14, 14, 14, 30])

    for a in apps:
        ws.append([
            a.get("rank"), a.get("company"), a.get("title"), a.get("location"),
            a.get("salary", "Not listed"), a.get("fit_score"), a.get("band"),
            a.get("reason"), a.get("cv_file"), a.get("url"),
            a.get("status", "Not applied"), a.get("date_applied", ""),
            a.get("follow_up_date", ""), a.get("notes", ""),
        ])
        row = ws.max_row
        fill = BAND.get(a.get("band"))
        if fill:
            ws.cell(row=row, column=7).fill = fill
        if a.get("url"):
            c = ws.cell(row=row, column=10)
            c.hyperlink = a["url"]
            c.value = "Open"
            c.font = Font(color="2563EB", underline="single")
        for col in range(1, 15):
            ws.cell(row=row, column=col).border = THIN
            ws.cell(row=row, column=col).alignment = Alignment(
                vertical="top", wrap_text=col in (8, 14))

    dv = DataValidation(
        type="list",
        formula1='"Not applied,Applied,Screening,Interview,Offer,Rejected,Withdrawn"',
        allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"K2:K{max(ws.max_row, 2)}")


def _cv_details(wb, apps):
    ws = wb.create_sheet("CV Details")
    _header(ws, ["Company", "Role", "CV File", "Keywords Matched",
                 "Keywords Missing", "[NEED NUMBER] Flags"],
            [22, 28, 34, 52, 42, 20])
    for a in apps:
        ws.append([
            a.get("company"), a.get("title"), a.get("cv_file"),
            ", ".join(a.get("keywords_matched", [])),
            ", ".join(a.get("keywords_missing", [])),
            a.get("need_number_flags", 0),
        ])
        for col in range(1, 7):
            ws.cell(row=ws.max_row, column=col).alignment = Alignment(
                vertical="top", wrap_text=col in (4, 5))


def _skipped(wb, skipped):
    ws = wb.create_sheet("Skipped Jobs")
    _header(ws, ["Company", "Role", "Location", "Fit", "Why skipped", "Job Link"],
            [22, 28, 18, 6, 56, 16])
    for s in skipped:
        ws.append([s.get("company"), s.get("title"), s.get("location"),
                   s.get("fit_score"), s.get("reason"), s.get("url", "")])
        ws.cell(row=ws.max_row, column=5).alignment = Alignment(wrap_text=True,
                                                               vertical="top")


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--applications", required=True)
    ap.add_argument("--out", required=True)
    args = ap.parse_args()

    with open(args.applications) as f:
        data = json.load(f)

    wb = Workbook()
    wb.remove(wb.active)
    _summary(wb, data)
    _applications(wb, data.get("applications", []))
    _cv_details(wb, data.get("applications", []))
    _skipped(wb, data.get("skipped", []))

    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    wb.save(args.out)
    print(f"Wrote {args.out}")


if __name__ == "__main__":
    main()
