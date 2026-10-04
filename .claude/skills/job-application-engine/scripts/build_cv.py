#!/usr/bin/env python3
"""Build an ATS-safe .docx CV with real typographic hierarchy.

Design constraint: everything here parses cleanly. Hierarchy comes from weight,
size, colour, spacing, paragraph borders and right tab stops — never from
columns, tables, text boxes, images, or headers/footers, which are the things
that actually scramble a parser.

Right-aligned dates use a right tab stop rather than a two-cell table. Same
visual result, no parse risk. That is the main trick in this file.

Usage:
    python3 build_cv.py --profile master_profile.json --tailor job_acme.json \
        --out "outputs/CVs/Jane_Doe_Acme_Ops_Manager.docx"
"""
import argparse
import json
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Inches, RGBColor

FONT = "Calibri"
BODY = 10.5
INK = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x5A, 0x5A, 0x5A)
RULE = "AAAAAA"
MARGIN = 0.65
CONTENT_W = 8.5 - (MARGIN * 2)


def _setup(doc):
    n = doc.styles["Normal"]
    n.font.name = FONT
    n.font.size = Pt(BODY)
    n.font.color.rgb = INK
    n.paragraph_format.space_after = Pt(0)
    n.paragraph_format.space_before = Pt(0)
    n.paragraph_format.line_spacing = 1.06
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(MARGIN)
        s.left_margin = s.right_margin = Inches(MARGIN)


def _p(doc, before=0, after=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    return p


def _run(p, text, size=BODY, bold=False, italic=False, color=INK, caps=False):
    r = p.add_run(text)
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    if caps:
        r.font.all_caps = True
    return r


def _rule(p, color=RULE, size=6):
    """Bottom border on a paragraph. Parsers ignore borders; humans see structure."""
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), color)
    bdr.append(bottom)
    pPr.append(bdr)


def _letterspace(run, twentieths=14):
    """Slight tracking on headings. Purely visual, text extracts unchanged."""
    rPr = run._element.get_or_add_rPr()
    sp = OxmlElement("w:spacing")
    sp.set(qn("w:val"), str(twentieths))
    rPr.append(sp)


def _right_tab(p):
    p.paragraph_format.tab_stops.add_tab_stop(
        Inches(CONTENT_W), WD_TAB_ALIGNMENT.RIGHT)


def _section(doc, title):
    p = _p(doc, before=11, after=4)
    r = _run(p, title, size=9.5, bold=True, caps=True)
    _letterspace(r)
    _rule(p)
    return p


def _bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    pf = p.paragraph_format
    pf.left_indent = Inches(0.20)
    pf.first_line_indent = Inches(-0.14)
    pf.space_before = Pt(0)
    pf.space_after = Pt(2.5)
    pf.line_spacing = 1.06
    _run(p, text)
    return p


def _dates(item):
    start = item.get("start", "")
    end = item.get("end") or "Present"
    return f"{start} – {end}".strip(" –")


def _contact(c):
    bits = [c.get(k) for k in ("location", "phone", "email", "linkedin", "portfolio")]
    return "   ·   ".join([b for b in bits if b])


def build(profile, tailor, out_path):
    doc = Document()
    _setup(doc)

    c = tailor.get("contact") or profile.get("contact", {})

    # Name — the one place we spend size
    p = _p(doc, after=2, align=WD_ALIGN_PARAGRAPH.CENTER)
    _letterspace(_run(p, c.get("name", ""), size=19, bold=True), 20)

    p = _p(doc, after=1, align=WD_ALIGN_PARAGRAPH.CENTER)
    _run(p, _contact(c), size=9, color=MUTED)
    if c.get("headline"):
        p = _p(doc, before=2, align=WD_ALIGN_PARAGRAPH.CENTER)
        _run(p, c["headline"], size=10, italic=True, color=MUTED)

    summary = tailor.get("summary") or profile.get("summary")
    if summary:
        _section(doc, "Summary")
        _run(_p(doc, after=1), summary)

    # Experience — title bold, company regular, dates right-tabbed
    for key, label in (("experience", "Experience"),
                       ("projects_detailed", "Projects")):
        roles = tailor.get(key) or profile.get(key) or []
        if not roles:
            continue
        _section(doc, label)
        for i, role in enumerate(roles):
            p = _p(doc, before=0 if i == 0 else 7)
            _right_tab(p)
            _run(p, role.get("title", ""), size=11, bold=True)
            if role.get("company"):
                _run(p, f"  ·  {role['company']}", size=11)
            if _dates(role):
                _run(p, "\t" + _dates(role), size=9.5, color=MUTED)
            sub_text = role.get("location") or role.get("context")
            if sub_text:
                _run(_p(doc, after=3), sub_text, size=9.5, italic=True, color=MUTED)
            else:
                p.paragraph_format.space_after = Pt(3)
            for b in role.get("bullets", []):
                _bullet(doc, b)

    edu = tailor.get("education") or profile.get("education", [])
    if edu:
        _section(doc, "Education")
        for i, e in enumerate(edu):
            p = _p(doc, before=0 if i == 0 else 5)
            _right_tab(p)
            _run(p, e.get("degree", ""), bold=True)
            if e.get("institution"):
                _run(p, f"  ·  {e['institution']}")
            if e.get("year"):
                _run(p, "\t" + str(e["year"]), size=9.5, color=MUTED)
            if e.get("detail"):
                _run(_p(doc, after=1), e["detail"], size=9.5, italic=True, color=MUTED)

    # Skills — label bold, values run-on. No rating bars, ever.
    skills = tailor.get("skills") or profile.get("skills", {})
    if skills:
        _section(doc, "Skills")
        if isinstance(skills, dict):
            for label, items in skills.items():
                if not items:
                    continue
                p = _p(doc, after=2)
                _run(p, f"{label.replace('_', ' ').title()}:  ", bold=True)
                _run(p, ", ".join(items))
        else:
            _run(_p(doc, after=2), ", ".join(skills))

    certs = tailor.get("certifications") or profile.get("certifications", [])
    if certs:
        _section(doc, "Certifications")
        for item in certs:
            _bullet(doc, item if isinstance(item, str) else item.get("name", ""))

    extras = tailor.get("projects") or profile.get("projects", [])
    has_detailed = tailor.get("projects_detailed") or profile.get("projects_detailed")
    if extras and not has_detailed:
        _section(doc, "Projects")
        for pr in extras:
            _bullet(doc, f"{pr.get('name', '')} — {pr.get('detail', '')}".strip(" —"))

    langs = tailor.get("languages") or profile.get("languages", [])
    if langs:
        _section(doc, "Languages")
        _run(_p(doc, after=1), ", ".join(langs))

    os.makedirs(os.path.dirname(os.path.abspath(out_path)), exist_ok=True)
    doc.save(out_path)
    return out_path


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--profile", required=True)
    ap.add_argument("--tailor", required=True)
    ap.add_argument("--out", required=True)
    a = ap.parse_args()
    profile = json.load(open(a.profile))
    tailor = json.load(open(a.tailor))
    path = build(profile, tailor, a.out)
    flags = json.dumps(tailor).count("[NEED NUMBER")
    print(f"Wrote {path}")
    if flags:
        print(f"WARNING: {flags} [NEED NUMBER] placeholder(s) unfilled.")


if __name__ == "__main__":
    main()
